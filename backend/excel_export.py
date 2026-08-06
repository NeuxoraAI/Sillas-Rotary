"""Shared Excel export builder for the Admin and Técnico beneficiarios exports.

Single source of truth for column layout and styling. Both `routers/admin.py`
and `routers/tecnica.py` call `build_beneficiarios_workbook` with the same row
shape so the two roles cannot drift into producing different files again (the
old, now-replaced implementation duplicated this logic inline in both files).

The workbook has 3 sheets, split by audience rather than by source table:
- "Beneficiario": who the person is, contact, address, tutores.
- "Información técnica": what the manufacturing team needs to build the chair.
- "Gestión y seguimiento": priority/entity/justification for case management.

Each row dict is expected to carry every key referenced below; callers build
it from a single SQL query (see `_ROW_QUERY_COLUMNS` docstring in the callers)
plus any resolved storage URLs.
"""

import io
from datetime import date, datetime
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

_HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F4E78")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
_DATE_NUMBER_FORMAT = "DD/MM/YYYY"
_MIN_COL_WIDTH = 10
_MAX_COL_WIDTH = 45


def _parse_fecha(value) -> Optional[date]:
    """Best-effort parse of the TEXT-stored fecha_* columns into a real date.

    Returns None (not a placeholder string) so openpyxl leaves the cell
    genuinely empty for missing dates instead of writing an empty string.
    """
    if not value:
        return None
    if isinstance(value, date):
        return value
    for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            continue
    return None


def _bool_label(value: Optional[bool]) -> str:
    if value is True:
        return "Sí"
    if value is False:
        return "No"
    return ""


def _nombre_completo(row: dict, prefix: str = "") -> str:
    """Prefer structured nombres/apellidos; fall back to the legacy compound
    `nombre` field for rows written before the structured columns existed."""
    nombres = row.get(f"{prefix}nombres")
    paterno = row.get(f"{prefix}apellido_paterno")
    materno = row.get(f"{prefix}apellido_materno")
    structured = " ".join(p for p in (nombres, paterno, materno) if p and str(p).strip())
    if structured.strip():
        return structured.strip()
    return (row.get(f"{prefix}nombre") or "").strip()


def _unidad_medida_mostrada(row: dict) -> str:
    """Measurements are stored verbatim in the capture unit while a solicitud
    is 'borrador', and converted to canonical inches only once it's
    'completo'. unidad_captura is a permanent audit trail of what was
    originally typed — it never flips to reflect that conversion — so it
    must NOT be read as "the unit of the stored value" for completed rows.
    Mirrors the identical rule in vista_tecnicos.html / admin-beneficiarios.html.
    """
    if row.get("solicitud_status") == "borrador":
        return row.get("unidad_captura") or "in"
    return "in"


def _unidad_peso_mostrada(row: dict) -> str:
    if row.get("solicitud_status") == "borrador":
        return row.get("unidad_peso_captura") or "lb"
    return "lb"


def _write_sheet(
    wb: Workbook,
    *,
    title: str,
    table_name: str,
    headers: list[str],
    rows: list[list],
    date_col_indices: tuple[int, ...] = (),
    hyperlink_col_indices: tuple[int, ...] = (),
    is_first: bool,
) -> None:
    ws = wb.active if is_first else wb.create_sheet()
    ws.title = title

    ws.append(headers)
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _HEADER_ALIGN
    ws.freeze_panes = "A2"

    for row in rows:
        ws.append(row)

    for row_idx, row in enumerate(rows, start=2):
        for col_idx in date_col_indices:
            cell = ws.cell(row=row_idx, column=col_idx)
            if cell.value is not None:
                cell.number_format = _DATE_NUMBER_FORMAT
        for col_idx in hyperlink_col_indices:
            cell = ws.cell(row=row_idx, column=col_idx)
            url = row[col_idx - 1]
            if url:
                cell.hyperlink = url
                cell.value = "Ver"
                cell.font = Font(color="1F4E78", underline="single")

    for idx, header in enumerate(headers, start=1):
        max_len = len(str(header))
        for row in rows:
            val = row[idx - 1]
            if val is not None and val != "":
                max_len = max(max_len, len(str(val)))
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = min(max(max_len + 2, _MIN_COL_WIDTH), _MAX_COL_WIDTH)

    if rows:
        last_col = get_column_letter(len(headers))
        table = Table(displayName=table_name, ref=f"A1:{last_col}{len(rows) + 1}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium9",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)


def build_beneficiarios_workbook(rows: list[dict]) -> io.BytesIO:
    """Build the 3-sheet beneficiarios export workbook from already-fetched rows.

    Each row dict must contain every key referenced in this function — see
    the SELECT in `routers/tecnica.py::exportar_beneficiarios_tecnica` (the
    query is identical for the admin export) for the exact source columns,
    plus `foto_url_resolved` (added by the caller via `_resolve_storage_url`).
    """
    wb = Workbook()

    # ── Hoja 1: Beneficiario ────────────────────────────────────────────
    beneficiario_headers = [
        "CURP", "Nombres", "Apellido paterno", "Apellido materno",
        "Fecha nacimiento", "Sexo", "Diagnóstico", "Teléfono",
        "Calle", "Núm. Ext.", "Núm. Int.", "Colonia", "Ciudad", "Estado",
        "País", "Región", "Sede",
        "Tutor 1 - Nombre", "Tutor 1 - Email",
        "Tutor 2 - Nombre", "Tutor 2 - Email",
        "Status",
    ]
    beneficiario_rows = []
    for row in rows:
        beneficiario_rows.append([
            row.get("curp_benef") or "",
            row.get("nombres") or "",
            row.get("apellido_paterno") or "",
            row.get("apellido_materno") or "",
            _parse_fecha(row.get("fecha_nacimiento")),
            row.get("sexo") or "",
            row.get("diagnostico") or "",
            row.get("telefonos") or "",
            row.get("calle") or "",
            row.get("num_ext") or "",
            row.get("num_int") or "",
            row.get("colonia") or "",
            row.get("ciudad") or "",
            row.get("estado_nombre") or "",
            row.get("pais_nombre") or "",
            row.get("region_nombre") or "",
            row.get("sede") or "",
            row.get("tutor1_nombre") or "",
            row.get("tutor1_email") or "",
            row.get("tutor2_nombre") or "",
            row.get("tutor2_email") or "",
            row.get("solicitud_status") or "",
        ])
    _write_sheet(
        wb,
        title="Beneficiario",
        table_name="Beneficiario",
        headers=beneficiario_headers,
        rows=beneficiario_rows,
        date_col_indices=(5,),
        is_first=True,
    )

    # ── Hoja 2: Información técnica ─────────────────────────────────────
    tecnica_headers = [
        "CURP", "Nombre", "Entorno",
        "Control tronco", "Control cabeza", "Control piernas",
        "Padecimiento", "Soporte oxígeno",
        "Altura total", "Unidad longitud", "Peso", "Unidad peso",
        "Cabeza-Asiento", "Hombro-Asiento", "Prof. Asiento",
        "Rodilla-Talón", "Ancho Cadera",
        "Equipo solicitado", "Foto",
    ]
    tecnica_rows = []
    for row in rows:
        unidad_medida = _unidad_medida_mostrada(row)
        unidad_peso = _unidad_peso_mostrada(row)
        tecnica_rows.append([
            row.get("curp_benef") or "",
            _nombre_completo(row),
            row.get("entorno") or "",
            row.get("control_tronco") or "",
            row.get("control_cabeza") or "",
            row.get("control_de_piernas") or "",
            row.get("padecimiento") or "",
            _bool_label(row.get("soporte_oxigeno")),
            row.get("altura_total_in"),
            unidad_medida,
            row.get("peso_kg"),
            unidad_peso,
            row.get("medida_cabeza_asiento"),
            row.get("medida_hombro_asiento"),
            row.get("medida_prof_asiento"),
            row.get("medida_rodilla_talon"),
            row.get("medida_ancho_cadera"),
            row.get("equipo_solicitado") or "",
            row.get("foto_url_resolved") or "",
        ])
    _write_sheet(
        wb,
        title="Información técnica",
        table_name="InformacionTecnica",
        headers=tecnica_headers,
        rows=tecnica_rows,
        hyperlink_col_indices=(19,),
        is_first=False,
    )

    # ── Hoja 3: Gestión y seguimiento ───────────────────────────────────
    gestion_headers = [
        "CURP", "Nombre", "Prioridad", "Entidad solicitante",
        "Justificación", "Fecha de estudio", "Status",
    ]
    gestion_rows = []
    for row in rows:
        gestion_rows.append([
            row.get("curp_benef") or "",
            _nombre_completo(row),
            row.get("prioridad") or "",
            row.get("entidad_solicitante") or "",
            row.get("justificacion") or "",
            _parse_fecha(row.get("fecha_estudio")),
            row.get("solicitud_status") or "",
        ])
    _write_sheet(
        wb,
        title="Gestión y seguimiento",
        table_name="GestionSeguimiento",
        headers=gestion_headers,
        rows=gestion_rows,
        date_col_indices=(6,),
        is_first=False,
    )

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
