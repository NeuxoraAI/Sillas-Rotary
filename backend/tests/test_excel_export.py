"""
Tests for the redesigned beneficiarios Excel export (admin + técnico).

Covers: 3-sheet structure, admin/técnico content parity, correct unit display
for borrador vs completo solicitudes, both tutores exported, real date cells,
and the specific bugs fixed in the redesign (swapped headers, dead columns,
missing equipo_solicitado/prioridad).
"""
import io

import openpyxl
import pytest


def _seed_beneficiario(
    _test_db_conn,
    region_lon,
    capturista_user,
    *,
    nombres="EXPORT",
    apellido_paterno="TEST",
    curp="EXPT000101HGTXXX01",
    tutor2=False,
    solicitud_status="completo",
    unidad_captura="in",
    unidad_peso_captura="lb",
    altura_total_in=25.0,
    peso_kg=110.0,
    prioridad="Alta",
    equipo_solicitado="Silla de ruedas neurológica PCI",
):
    with _test_db_conn.cursor() as cur:
        cur.execute(
            """
            INSERT INTO beneficiarios
                (nombre, nombres, apellido_paterno, apellido_materno, curp_benef,
                 fecha_nacimiento, sexo, diagnostico, calle, num_ext, num_int,
                 colonia, ciudad, estado_codigo, estado_nombre, telefonos,
                 region_id, sede)
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
            RETURNING id
            """,
            (
                f"{nombres} {apellido_paterno} MUESTRA", nombres, apellido_paterno, "MUESTRA", curp,
                "2015-06-01", "M", "Parálisis cerebral", "Calle Uno", "10", "2",
                "Centro", "León", "11", "GUANAJUATO", "4621234567",
                region_lon["id"], "León sede Forum",
            ),
        )
        beneficiario_id = cur.fetchone()[0]

        cur.execute(
            """
            INSERT INTO tutores
                (beneficiario_id, numero_tutor, nombres, apellido_paterno, apellido_materno, email)
            VALUES (%s, 1, %s, %s, %s, %s)
            """,
            (beneficiario_id, "TUTOR", "UNO", "MUESTRA", "tutor1@test.mx"),
        )
        if tutor2:
            cur.execute(
                """
                INSERT INTO tutores
                    (beneficiario_id, numero_tutor, nombres, apellido_paterno, apellido_materno, email)
                VALUES (%s, 2, %s, %s, %s, %s)
                """,
                (beneficiario_id, "TUTOR", "DOS", "MUESTRA", "tutor2@test.mx"),
            )

        cur.execute(
            """
            INSERT INTO estudios_socioeconomicos
                (beneficiario_id, usuario_id, sede, fecha_estudio, status)
            VALUES (%s, %s, %s, %s, 'completo')
            """,
            (beneficiario_id, capturista_user["id"], "León sede Forum", "2026-04-18"),
        )

        cur.execute(
            """
            INSERT INTO solicitudes_tecnicas
                (beneficiario_id, usuario_id, entorno, control_tronco, control_cabeza,
                 control_de_piernas, padecimiento, soporte_oxigeno,
                 altura_total_in, peso_kg, medida_cabeza_asiento, medida_hombro_asiento,
                 medida_prof_asiento, medida_rodilla_talon, medida_ancho_cadera,
                 unidad_captura, unidad_peso_captura, equipo_solicitado,
                 entidad_solicitante, prioridad, justificacion, status)
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
            """,
            (
                beneficiario_id, capturista_user["id"], "Urbano / Interiores", "Completo", "Independiente",
                "Parcial", "Escoliosis", True,
                altura_total_in, peso_kg, 10.0, 12.0,
                14.0, 16.0, 18.0,
                unidad_captura, unidad_peso_captura, equipo_solicitado,
                "Rotary Club León", prioridad, "Caso urgente", solicitud_status,
            ),
        )
    _test_db_conn.commit()
    return beneficiario_id


def _load_workbook(response) -> openpyxl.Workbook:
    return openpyxl.load_workbook(io.BytesIO(response.content), data_only=True)


class TestExcelExportStructure:
    def test_admin_export_has_three_sheets(self, client, admin_headers, _test_db_conn, region_lon, capturista_user):
        _seed_beneficiario(_test_db_conn, region_lon, capturista_user, curp="EXPT000101HGTXXX02")
        res = client.get("/api/admin/beneficiarios/export", headers=admin_headers)
        assert res.status_code == 200
        wb = _load_workbook(res)
        assert wb.sheetnames == ["Beneficiario", "Información técnica", "Gestión y seguimiento"]

    def test_tecnico_export_has_same_three_sheets(self, client, tecnico_headers, _test_db_conn, region_lon, capturista_user):
        _seed_beneficiario(_test_db_conn, region_lon, capturista_user, curp="EXPT000101HGTXXX03")
        res = client.get("/api/tecnica/beneficiarios/export", headers=tecnico_headers)
        assert res.status_code == 200
        wb = _load_workbook(res)
        assert wb.sheetnames == ["Beneficiario", "Información técnica", "Gestión y seguimiento"]

    def test_no_dead_placeholder_columns(self, client, admin_headers, _test_db_conn, region_lon, capturista_user):
        """Regression guard: the old export had 2 columns hardcoded to "" with
        no backing DB field ("teléfono adicional", "QUIEN CANALIZA"). Neither
        header should exist anymore."""
        _seed_beneficiario(_test_db_conn, region_lon, capturista_user, curp="EXPT000101HGTXXX04")
        res = client.get("/api/admin/beneficiarios/export", headers=admin_headers)
        wb = _load_workbook(res)
        all_headers = set()
        for ws in wb.worksheets:
            all_headers.update(c.value for c in ws[1])
        assert "No. de teléfono adicional" not in all_headers
        assert "QUIEN CANALIZA" not in all_headers


class TestExcelExportContent:
    def test_equipo_solicitado_and_prioridad_present(self, client, admin_headers, _test_db_conn, region_lon, capturista_user):
        """Regression guard: these were the two most operationally important
        fields missing from the old export."""
        _seed_beneficiario(
            _test_db_conn, region_lon, capturista_user, curp="EXPT000101HGTXXX05",
            equipo_solicitado="Silla de ruedas neurológica PCA", prioridad="Alta",
        )
        res = client.get("/api/admin/beneficiarios/export", headers=admin_headers)
        wb = _load_workbook(res)

        tecnica_ws = wb["Información técnica"]
        headers = [c.value for c in tecnica_ws[1]]
        equipo_col = headers.index("Equipo solicitado")
        row = list(tecnica_ws.iter_rows(min_row=2, max_row=2, values_only=True))[0]
        assert row[equipo_col] == "Silla de ruedas neurológica PCA"

        gestion_ws = wb["Gestión y seguimiento"]
        headers = [c.value for c in gestion_ws[1]]
        prioridad_col = headers.index("Prioridad")
        row = list(gestion_ws.iter_rows(min_row=2, max_row=2, values_only=True))[0]
        assert row[prioridad_col] == "Alta"

    def test_padecimiento_and_diagnostico_not_swapped(self, client, admin_headers, _test_db_conn, region_lon, capturista_user):
        """Regression guard: the old export's "Padecimiento" header actually
        showed diagnostico, and vice versa. Confirm each field lands under
        its own correctly-named column now."""
        beneficiario_id = _seed_beneficiario(_test_db_conn, region_lon, capturista_user, curp="EXPT000101HGTXXX06")
        res = client.get("/api/admin/beneficiarios/export", headers=admin_headers)
        wb = _load_workbook(res)

        ben_ws = wb["Beneficiario"]
        ben_headers = [c.value for c in ben_ws[1]]
        diag_col = ben_headers.index("Diagnóstico")
        ben_row = list(ben_ws.iter_rows(min_row=2, max_row=2, values_only=True))[0]
        assert ben_row[diag_col] == "Parálisis cerebral"

        tecnica_ws = wb["Información técnica"]
        tec_headers = [c.value for c in tecnica_ws[1]]
        padecimiento_col = tec_headers.index("Padecimiento")
        tec_row = list(tecnica_ws.iter_rows(min_row=2, max_row=2, values_only=True))[0]
        assert tec_row[padecimiento_col] == "Escoliosis"

    def test_both_tutores_exported(self, client, admin_headers, _test_db_conn, region_lon, capturista_user):
        """Regression guard: the old export only ever included tutor #1."""
        _seed_beneficiario(_test_db_conn, region_lon, capturista_user, curp="EXPT000101HGTXXX07", tutor2=True)
        res = client.get("/api/admin/beneficiarios/export", headers=admin_headers)
        wb = _load_workbook(res)
        ws = wb["Beneficiario"]
        headers = [c.value for c in ws[1]]
        t1_col = headers.index("Tutor 1 - Nombre")
        t2_col = headers.index("Tutor 2 - Nombre")
        row = list(ws.iter_rows(min_row=2, max_row=2, values_only=True))[0]
        assert "TUTOR UNO" in row[t1_col]
        assert "TUTOR DOS" in row[t2_col]

    def test_single_tutor_leaves_tutor2_blank_not_missing(self, client, admin_headers, _test_db_conn, region_lon, capturista_user):
        _seed_beneficiario(_test_db_conn, region_lon, capturista_user, curp="EXPT000101HGTXXX08", tutor2=False)
        res = client.get("/api/admin/beneficiarios/export", headers=admin_headers)
        wb = _load_workbook(res)
        ws = wb["Beneficiario"]
        headers = [c.value for c in ws[1]]
        t2_col = headers.index("Tutor 2 - Nombre")
        row = list(ws.iter_rows(min_row=2, max_row=2, values_only=True))[0]
        assert row[t2_col] in (None, "")

    def test_fecha_nacimiento_is_real_date_cell(self, client, admin_headers, _test_db_conn, region_lon, capturista_user):
        """Regression guard: dates used to be written as raw TEXT with no
        number_format; confirm the redesign writes real date-typed cells."""
        import datetime
        _seed_beneficiario(_test_db_conn, region_lon, capturista_user, curp="EXPT000101HGTXXX09")
        res = client.get("/api/admin/beneficiarios/export", headers=admin_headers)
        wb = _load_workbook(res)
        ws = wb["Beneficiario"]
        headers = [c.value for c in ws[1]]
        fecha_col_idx = headers.index("Fecha nacimiento") + 1
        cell = ws.cell(row=2, column=fecha_col_idx)
        assert isinstance(cell.value, (datetime.date, datetime.datetime))


class TestExcelExportUnits:
    """The most delicate part of the redesign: unidad_captura/unidad_peso_captura
    are a permanent audit trail that never flips to "in"/"lb" even after the
    stored value is converted to canonical at finalize time. The export must
    only trust that column for borrador rows."""

    def test_borrador_shows_captured_unit_not_canonical(self, client, admin_headers, _test_db_conn, region_lon, capturista_user):
        _seed_beneficiario(
            _test_db_conn, region_lon, capturista_user, curp="EXPT000101HGTXXX10",
            solicitud_status="borrador", unidad_captura="cm", unidad_peso_captura="kg",
            altura_total_in=63.5, peso_kg=50.0,
        )
        res = client.get("/api/admin/beneficiarios/export", headers=admin_headers)
        wb = _load_workbook(res)
        ws = wb["Información técnica"]
        headers = [c.value for c in ws[1]]
        unidad_long_col = headers.index("Unidad longitud")
        unidad_peso_col = headers.index("Unidad peso")
        row = list(ws.iter_rows(min_row=2, max_row=2, values_only=True))[0]
        assert row[unidad_long_col] == "cm"
        assert row[unidad_peso_col] == "kg"

    def test_completo_shows_canonical_unit_regardless_of_captured_unit(self, client, admin_headers, _test_db_conn, region_lon, capturista_user):
        """A completo solicitud's unidad_captura/unidad_peso_captura still say
        whatever was originally typed (e.g. "cm") — the export must show the
        canonical in/lb unit for these rows, not the stale audit-trail value."""
        _seed_beneficiario(
            _test_db_conn, region_lon, capturista_user, curp="EXPT000101HGTXXX11",
            solicitud_status="completo", unidad_captura="cm", unidad_peso_captura="kg",
            altura_total_in=25.0, peso_kg=110.0,
        )
        res = client.get("/api/admin/beneficiarios/export", headers=admin_headers)
        wb = _load_workbook(res)
        ws = wb["Información técnica"]
        headers = [c.value for c in ws[1]]
        unidad_long_col = headers.index("Unidad longitud")
        unidad_peso_col = headers.index("Unidad peso")
        row = list(ws.iter_rows(min_row=2, max_row=2, values_only=True))[0]
        assert row[unidad_long_col] == "in"
        assert row[unidad_peso_col] == "lb"


class TestExcelExportAdminTecnicoParity:
    def test_same_filter_produces_same_cell_content(self, client, admin_headers, tecnico_headers, _test_db_conn, region_lon, capturista_user):
        """Admin and técnico must produce the same information for the same
        filter — the shared excel_export.build_beneficiarios_workbook module
        is what guarantees this mechanically."""
        _seed_beneficiario(_test_db_conn, region_lon, capturista_user, curp="EXPT000101HGTXXX12", tutor2=True)

        admin_res = client.get("/api/admin/beneficiarios/export?q=EXPT000101HGTXXX12", headers=admin_headers)
        tecnico_res = client.get("/api/tecnica/beneficiarios/export?q=EXPT000101HGTXXX12", headers=tecnico_headers)
        assert admin_res.status_code == 200
        assert tecnico_res.status_code == 200

        admin_wb = _load_workbook(admin_res)
        tecnico_wb = _load_workbook(tecnico_res)
        assert admin_wb.sheetnames == tecnico_wb.sheetnames
        for name in admin_wb.sheetnames:
            admin_rows = list(admin_wb[name].iter_rows(values_only=True))
            tecnico_rows = list(tecnico_wb[name].iter_rows(values_only=True))
            assert admin_rows == tecnico_rows, f"Mismatch in sheet {name}"


class TestExcelExportFilters:
    """Confirm the redesign didn't change WHERE-clause behavior — only the
    SELECT list grew, filters must still narrow results the same way."""

    def test_q_filter_still_narrows_results(self, client, admin_headers, _test_db_conn, region_lon, capturista_user):
        _seed_beneficiario(_test_db_conn, region_lon, capturista_user, curp="EXPT000101HGTXXX13", nombres="FILTROUNO")
        _seed_beneficiario(_test_db_conn, region_lon, capturista_user, curp="EXPT000101HGTXXX14", nombres="FILTRODOS")

        res = client.get("/api/admin/beneficiarios/export?q=FILTROUNO", headers=admin_headers)
        wb = _load_workbook(res)
        ws = wb["Beneficiario"]
        names = [row[1] for row in ws.iter_rows(min_row=2, values_only=True)]
        assert any("FILTROUNO" in (n or "") for n in names)
        assert not any("FILTRODOS" in (n or "") for n in names)

    def test_ids_filter_still_works(self, client, admin_headers, _test_db_conn, region_lon, capturista_user):
        beneficiario_id = _seed_beneficiario(_test_db_conn, region_lon, capturista_user, curp="EXPT000101HGTXXX15")
        res = client.get(f"/api/admin/beneficiarios/export?ids={beneficiario_id}", headers=admin_headers)
        assert res.status_code == 200
        wb = _load_workbook(res)
        ws = wb["Beneficiario"]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        assert len(rows) == 1


class TestExcelExportDuplicateRows:
    def test_beneficiario_with_multiple_solicitudes_exports_once(
        self, client, admin_headers, _test_db_conn, region_lon, capturista_user, admin_user,
    ):
        """A beneficiario can have more than one estudio/solicitud if captured
        by different usuario_id (UNIQUE constraint is per beneficiario+usuario,
        not beneficiario alone — migration 0009). The export must only ever
        show the most recent one per beneficiario, not fan out into duplicate
        rows."""
        beneficiario_id = _seed_beneficiario(_test_db_conn, region_lon, capturista_user, curp="EXPT000101HGTXXX16")
        with _test_db_conn.cursor() as cur:
            cur.execute(
                """
                INSERT INTO estudios_socioeconomicos
                    (beneficiario_id, usuario_id, sede, fecha_estudio, status)
                VALUES (%s, %s, %s, %s, 'completo')
                """,
                (beneficiario_id, admin_user["id"], "Otra sede", "2026-05-01"),
            )
            cur.execute(
                """
                INSERT INTO solicitudes_tecnicas
                    (beneficiario_id, usuario_id, entorno, control_tronco, control_cabeza,
                     control_de_piernas, altura_total_in, peso_kg, status)
                VALUES (%s, %s, 'Urbano / Interiores', 'Completo', 'Independiente', 'Parcial', 30.0, 120.0, 'completo')
                """,
                (beneficiario_id, admin_user["id"]),
            )
        _test_db_conn.commit()

        res = client.get(f"/api/admin/beneficiarios/export?ids={beneficiario_id}", headers=admin_headers)
        wb = _load_workbook(res)
        ws = wb["Beneficiario"]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        assert len(rows) == 1, "beneficiario with 2 estudios/solicitudes should export exactly one row (most recent)"
